"""Small helper functions used by app.py (no Streamlit code in this file)."""
import textwrap
from io import BytesIO

import pandas as pd
from docx import Document
from reportlab.pdfgen import canvas


def find_col(columns, keywords):
    """Return the first column whose name contains one of the keywords."""
    for keyword in keywords:
        for col in columns:
            if keyword in col.lower():
                return col
    return None


def fmt(value):
    """Format a number for display; None becomes N/A."""
    if value is None or pd.isna(value):
        return "N/A"
    return f"{value:,.2f}"


def to_number(series):
    """Convert to numbers; text like '₹1,200' or '$ 1,200' becomes 1200."""
    if not pd.api.types.is_numeric_dtype(series):
        cleaned = series.astype(str).str.replace(r"[₹$,\s]", "", regex=True)
        series = cleaned.where(series.notna() & ~cleaned.isin(["", "nan", "None"]))
    return pd.to_numeric(series, errors="coerce")


def to_date(series, dayfirst=False):
    """Convert to dates; dayfirst=True reads 05/03/2026 as 5 March."""
    return pd.to_datetime(series, errors="coerce", dayfirst=dayfirst)


def health_check(data, sales, profit, quantity, date, dayfirst=False):
    """Return (score, issues) describing the quality of the raw data."""
    total_rows = len(data)
    bad_rows = pd.Series(False, index=data.index)

    missing = data.isna()
    duplicates = data.duplicated()
    bad_rows |= missing.any(axis=1)
    bad_rows |= duplicates

    not_numeric = pd.Series(False, index=data.index)
    negative = pd.Series(False, index=data.index)

    for col in [sales, profit, quantity]:
        if col:
            converted = to_number(data[col])
            not_numeric |= converted.isna() & data[col].notna()
            if col != profit:  # negative profit (loss) is valid
                negative |= converted < 0

    invalid_dates = pd.Series(False, index=data.index)
    if date:
        parsed = to_date(data[date], dayfirst)
        invalid_dates = parsed.isna() & data[date].notna()

    bad_rows |= not_numeric | negative | invalid_dates

    issues = {
        "Missing values": int(missing.sum().sum()),
        "Duplicate rows": int(duplicates.sum()),
        "Non-numeric values in number columns": int(not_numeric.sum()),
        "Negative sales / quantity values": int(negative.sum()),
        "Invalid dates": int(invalid_dates.sum()),
    }

    score = 100.0 * (1 - bad_rows.sum() / total_rows) if total_rows else 0.0
    return score, issues


def partial_month_note(month, last_date):
    """Warn when a month is the final month of the data and not fully covered."""
    if (
        last_date is not None
        and not pd.isna(last_date)
        and month == last_date.strftime("%Y-%m")
        and last_date.day < last_date.days_in_month
    ):
        return f" (Note: {month} may be incomplete, data ends on {last_date.date()}.)"
    return ""


def group_table(data, dim, sales, profit):
    """Sales, profit and margin % grouped by a column."""
    value_cols = [c for c in [sales, profit] if c]
    table = data.groupby(dim)[value_cols].sum()
    if sales and profit:
        table["Margin %"] = (
            table[profit] / table[sales].replace(0, float("nan")) * 100
        ).round(2)
    return table.reset_index()


def build_docx_report(report_lines):
    """Return the business report as an in-memory Word file."""
    doc = Document()
    doc.add_heading("Hybrid Data Analytics & AI Reporting Engine", level=1)
    doc.add_heading("AI Business Report", level=2)
    for line in report_lines:
        doc.add_paragraph(line)
    output = BytesIO()
    doc.save(output)
    output.seek(0)
    return output


def build_pdf_report(report_lines):
    """Return the business report as an in-memory PDF file."""
    output = BytesIO()
    pdf = canvas.Canvas(output)
    pdf.setTitle("AI Business Report")
    pdf.drawString(50, 800, "Hybrid Data Analytics & AI Reporting Engine")
    pdf.drawString(50, 775, "AI Business Report")

    y = 740
    for line in report_lines:
        for part in textwrap.wrap(line, 90) or [""]:
            if y < 60:
                pdf.showPage()
                y = 800
            pdf.drawString(50, y, part)
            y -= 18
        y -= 8

    pdf.save()
    output.seek(0)
    return output
